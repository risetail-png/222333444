"""
P8セル 写真差し替えツール
============================================================
エクセルファイル（橋梁点検票）のP8セルに埋め込まれた写真を、
指定した写真フォルダ内の矢印記号付きJPGに自動差し替えするツール。

使い方:
    python p8_photo_replacer.py

必要ライブラリ:
    pip install -r requirements.txt
"""

import json
import sys
import os
import threading
import unicodedata
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage

# ── 定数 ────────────────────────────────────────────────────
ARROW_SYMBOLS = ["→", "←", "↑", "↓", "↗", "↘", "↙", "↖", "⇒", "⇐", "⇑", "⇓"]
P8_COL = 15   # P列 (0-indexed)
P8_ROW = 7    # 8行目 (0-indexed)
CONFIG_FILE = Path.home() / ".p8_photo_replacer_config.json"

# ── カラーテーマ ─────────────────────────────────────────────
BG       = "#1a1d23"
BG2      = "#22262f"
BG3      = "#2b3040"
ACCENT   = "#4f8ef7"
ACCENT2  = "#6ba3ff"
SUCCESS  = "#3ecf8e"
ERROR    = "#f76f6f"
WARNING  = "#f7b84f"
TEXT     = "#e8eaf0"
TEXT_DIM = "#8891a8"
BORDER   = "#353c50"


# ════════════════════════════════════════════════════════════
# コアロジック
# ════════════════════════════════════════════════════════════

def normalize(text: str) -> str:
    """全角→半角に統一し、小文字化・空白除去して返す"""
    return unicodedata.normalize("NFKC", text).lower().strip()


def get_bridge_name(excel_path: Path) -> str:
    """エクセルファイル名からアンダースコア区切りの末尾を橋梁名として返す"""
    parts = excel_path.stem.split("_")
    return parts[-1] if len(parts) > 1 else excel_path.stem


def find_bridge_folder(photo_root: Path, bridge_name: str) -> Path | None:
    """橋梁名と部分一致（全角半角無視）するフォルダを返す"""
    norm = normalize(bridge_name)
    for folder in photo_root.iterdir():
        if folder.is_dir():
            f = normalize(folder.name)
            if norm in f or f in norm:
                return folder
    return None


def find_arrow_jpg(folder: Path) -> Path | None:
    """フォルダ内から矢印記号を含む最初のJPGファイルを返す"""
    for f in sorted(folder.iterdir()):
        if f.suffix.lower() in (".jpg", ".jpeg"):
            if any(sym in f.name for sym in ARROW_SYMBOLS):
                return f
    return None


def replace_p8_image(excel_path: Path, new_image_path: Path) -> tuple[bool, str]:
    """P8セルの画像を差し替えて「元ファイル名_更新済.xlsx」で保存する"""
    try:
        wb = load_workbook(str(excel_path))
    except Exception as e:
        return False, f"エクセル読み込みエラー: {e}"

    ws = wb["状態把握"] if "状態把握" in wb.sheetnames else wb.active

    target_idx = None
    for i, img in enumerate(ws._images):
        try:
            if img.anchor._from.col == P8_COL and img.anchor._from.row == P8_ROW:
                target_idx = i
                break
        except Exception:
            continue

    if target_idx is None:
        return False, "P8セルに画像が見つかりません"

    old_anchor = ws._images[target_idx].anchor
    new_img = XLImage(str(new_image_path))
    new_img.anchor = old_anchor
    ws._images[target_idx] = new_img

    out = excel_path.parent / f"{excel_path.stem}_更新済{excel_path.suffix}"
    try:
        wb.save(str(out))
    except Exception as e:
        return False, f"保存エラー: {e}"

    return True, str(out)


def run_batch(excel_folder: str, photo_folder: str, callback=None) -> dict:
    """フォルダ内の全エクセルを一括処理する"""
    results = {"ok": [], "ng": [], "total": 0}
    excel_files = sorted(Path(excel_folder).glob("*.xlsx"))
    results["total"] = len(excel_files)

    for excel_path in excel_files:
        bridge_name = get_bridge_name(excel_path)
        if callback:
            callback("info", excel_path.name, f"橋梁名: {bridge_name}")

        bridge_folder = find_bridge_folder(Path(photo_folder), bridge_name)
        if not bridge_folder:
            msg = "写真フォルダが見つかりません"
            results["ng"].append((excel_path.name, msg))
            if callback:
                callback("error", excel_path.name, msg)
            continue

        arrow_jpg = find_arrow_jpg(bridge_folder)
        if not arrow_jpg:
            msg = "矢印記号付きJPGが見つかりません"
            results["ng"].append((excel_path.name, msg))
            if callback:
                callback("error", excel_path.name, msg)
            continue

        if callback:
            callback("info", excel_path.name, f"使用画像: {arrow_jpg.name}")

        success, result = replace_p8_image(excel_path, arrow_jpg)
        if success:
            results["ok"].append((excel_path.name, result))
            if callback:
                callback("ok", excel_path.name, f"保存: {Path(result).name}")
        else:
            results["ng"].append((excel_path.name, result))
            if callback:
                callback("error", excel_path.name, result)

    return results


# ════════════════════════════════════════════════════════════
# GUI
# ════════════════════════════════════════════════════════════

def load_config() -> dict:
    try:
        if CONFIG_FILE.exists():
            return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
    except Exception:
        pass
    return {"excel_folder": "", "photo_folder": ""}


def save_config(data: dict):
    try:
        CONFIG_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("P8 写真差し替えツール")
        self.geometry("720x580")
        self.minsize(600, 480)
        self.configure(bg=BG)
        self._config = load_config()
        self._running = False
        self._build_ui()
        self._load_fields()

    def _build_ui(self):
        # タイトルバー
        title_bar = tk.Frame(self, bg=BG3, height=56)
        title_bar.pack(fill="x")
        title_bar.pack_propagate(False)
        tk.Label(title_bar, text="🗂  P8 写真差し替えツール",
                 font=("Yu Gothic UI", 15, "bold"),
                 bg=BG3, fg=TEXT, padx=20).pack(side="left", fill="y")
        tk.Label(title_bar, text="橋梁点検エクセル  写真自動更新",
                 font=("Yu Gothic UI", 9), bg=BG3, fg=TEXT_DIM, padx=10).pack(side="right", fill="y")

        main = tk.Frame(self, bg=BG, padx=24, pady=20)
        main.pack(fill="both", expand=True)

        # フォルダ設定
        tk.Label(main, text="フォルダ設定", font=("Yu Gothic UI", 10, "bold"),
                 bg=BG, fg=ACCENT2).pack(anchor="w", pady=(0, 10))

        self._excel_var = tk.StringVar()
        self._photo_var = tk.StringVar()
        self._folder_row(main, "📁  エクセルフォルダ", self._excel_var,
                         lambda: self._browse(self._excel_var))
        self._folder_row(main, "🖼  写真フォルダ    ", self._photo_var,
                         lambda: self._browse(self._photo_var))

        tk.Frame(main, bg=BORDER, height=1).pack(fill="x", pady=16)

        # 実行ボタン
        run_row = tk.Frame(main, bg=BG)
        run_row.pack(fill="x", pady=(0, 16))
        self._run_btn = tk.Button(
            run_row, text="▶  一括処理を開始", command=self._start,
            bg=ACCENT, fg="white", activebackground=ACCENT2, activeforeground="white",
            relief="flat", font=("Yu Gothic UI", 11, "bold"),
            padx=28, pady=10, cursor="hand2", bd=0)
        self._run_btn.pack(side="left")
        self._progress = ttk.Progressbar(run_row, mode="indeterminate", length=200)
        self._progress.pack(side="left", padx=20)
        self._progress_label = tk.Label(run_row, text="", font=("Yu Gothic UI", 9),
                                        bg=BG, fg=TEXT_DIM)
        self._progress_label.pack(side="left")
        style = ttk.Style()
        style.theme_use("default")
        style.configure("TProgressbar", troughcolor=BG2, background=ACCENT, thickness=8)

        # ログ
        tk.Label(main, text="処理ログ", font=("Yu Gothic UI", 10, "bold"),
                 bg=BG, fg=ACCENT2).pack(anchor="w", pady=(0, 6))
        log_frame = tk.Frame(main, bg=BORDER, padx=1, pady=1)
        log_frame.pack(fill="both", expand=True)
        inner = tk.Frame(log_frame, bg=BG2)
        inner.pack(fill="both", expand=True)
        self._log = tk.Text(inner, bg=BG2, fg=TEXT, insertbackground=TEXT,
                            relief="flat", font=("Consolas", 9),
                            padx=12, pady=10, state="disabled", wrap="word",
                            selectbackground=ACCENT)
        self._log.pack(side="left", fill="both", expand=True)
        sb = tk.Scrollbar(inner, command=self._log.yview, bg=BG2, troughcolor=BG2,
                          activebackground=ACCENT)
        sb.pack(side="right", fill="y")
        self._log.configure(yscrollcommand=sb.set)
        self._log.tag_configure("ok",      foreground=SUCCESS)
        self._log.tag_configure("error",   foreground=ERROR)
        self._log.tag_configure("info",    foreground=TEXT_DIM)
        self._log.tag_configure("header",  foreground=ACCENT2, font=("Consolas", 9, "bold"))
        self._log.tag_configure("summary", foreground=WARNING,  font=("Consolas", 9, "bold"))

        # ステータスバー
        bar = tk.Frame(self, bg=BG3, height=28)
        bar.pack(fill="x", side="bottom")
        bar.pack_propagate(False)
        self._status = tk.Label(bar, text="準備完了", font=("Yu Gothic UI", 8),
                                bg=BG3, fg=TEXT_DIM, padx=14)
        self._status.pack(side="left", fill="y")

    def _folder_row(self, parent, label, var, cmd):
        row = tk.Frame(parent, bg=BG, pady=5)
        row.pack(fill="x")
        tk.Label(row, text=label, font=("Yu Gothic UI", 9),
                 bg=BG, fg=TEXT_DIM, width=18, anchor="w").pack(side="left")
        tk.Entry(row, textvariable=var, bg=BG2, fg=TEXT, insertbackground=TEXT,
                 relief="flat", font=("Yu Gothic UI", 9),
                 highlightthickness=1, highlightcolor=ACCENT,
                 highlightbackground=BORDER).pack(side="left", fill="x", expand=True, ipady=6, padx=(0, 8))
        tk.Button(row, text="選択", command=cmd, bg=BG3, fg=TEXT,
                  activebackground=ACCENT, activeforeground="white",
                  relief="flat", font=("Yu Gothic UI", 9),
                  padx=14, pady=4, cursor="hand2", bd=0).pack(side="right")

    def _browse(self, var):
        path = filedialog.askdirectory()
        if path:
            var.set(path)
            self._save_fields()

    def _load_fields(self):
        self._excel_var.set(self._config.get("excel_folder", ""))
        self._photo_var.set(self._config.get("photo_folder", ""))

    def _save_fields(self):
        self._config["excel_folder"] = self._excel_var.get()
        self._config["photo_folder"] = self._photo_var.get()
        save_config(self._config)

    def _log_write(self, level, filename, message):
        self._log.configure(state="normal")
        prefix = {"ok": "✅", "error": "❌", "info": "   "}.get(level, "  ")
        self._log.insert("end", f"{prefix}  {filename}\n       {message}\n", level)
        self._log.see("end")
        self._log.configure(state="disabled")

    def _log_header(self, text):
        self._log.configure(state="normal")
        self._log.insert("end", f"\n{'─'*50}\n{text}\n{'─'*50}\n", "header")
        self._log.see("end")
        self._log.configure(state="disabled")

    def _log_summary(self, text):
        self._log.configure(state="normal")
        self._log.insert("end", f"\n{text}\n", "summary")
        self._log.see("end")
        self._log.configure(state="disabled")

    def _start(self):
        excel = self._excel_var.get().strip()
        photo = self._photo_var.get().strip()
        if not excel or not photo:
            messagebox.showwarning("入力エラー", "エクセルフォルダと写真フォルダを両方選択してください。")
            return
        if not Path(excel).exists():
            messagebox.showerror("エラー", f"エクセルフォルダが存在しません:\n{excel}")
            return
        if not Path(photo).exists():
            messagebox.showerror("エラー", f"写真フォルダが存在しません:\n{photo}")
            return
        self._save_fields()
        self._run_btn.configure(state="disabled", bg=BG3)
        self._progress.start(12)
        self._status.configure(text="処理中...")
        threading.Thread(target=self._run_thread, args=(excel, photo), daemon=True).start()

    def _run_thread(self, excel, photo):
        self.after(0, self._log_header, f"処理開始\nエクセル: {excel}\n写真:     {photo}")

        def cb(level, filename, message):
            self.after(0, self._log_write, level, filename, message)

        results = run_batch(excel, photo, callback=cb)
        ok, ng, total = len(results["ok"]), len(results["ng"]), results["total"]
        summary = f"完了  {ok}/{total} 件成功" + (f"  /  失敗 {ng} 件" if ng else "  ✅ 全件成功")
        self.after(0, self._log_summary, summary)
        self.after(0, self._finish, ok, ng, total)

    def _finish(self, ok, ng, total):
        self._progress.stop()
        self._run_btn.configure(state="normal", bg=ACCENT)
        self._status.configure(text=f"完了: {ok}/{total} 件成功" + (f"  失敗: {ng} 件" if ng else ""))
        if ng == 0:
            messagebox.showinfo("完了", f"{ok} 件すべて正常に処理しました。")
        else:
            messagebox.showwarning("完了（一部失敗）",
                                   f"成功: {ok} 件 / 失敗: {ng} 件\n詳細はログを確認してください。")


# ════════════════════════════════════════════════════════════
# エントリーポイント
# ════════════════════════════════════════════════════════════

if __name__ == "__main__":
    App().mainloop()
